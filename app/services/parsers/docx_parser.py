"""DOCX Parser converting Word documents to Markdown with table and merged-cell support."""

import io
import re
from typing import Any, Dict, List, Optional
from loguru import logger

try:
    import docx
    from docx.oxml.ns import qn
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    docx = None

from app.core.text_normalizer import normalize_text
from app.services.parsers.base import BaseParser, image_placeholder
from app.services.parsers.pdf_parser import enhance_vietnamese_headings


class _Numbering:
    """Render Word's stored list labels, which Paragraph.text does not include."""

    def __init__(self, document: Any):
        try:
            self.root = document.part.part_related_by(RT.NUMBERING).element
        except KeyError:
            self.root = None
        self.counters: Dict[str, Dict[int, int]] = {}

    @staticmethod
    def _value(element: Any, path: str, default: str = "") -> str:
        node = element.find(path) if element is not None else None
        return node.get(qn("w:val"), default) if node is not None else default

    def _level(self, num: Any, level: int) -> Any:
        override = num.find(f'w:lvlOverride[@w:ilvl="{level}"]', num.nsmap)
        if override is not None and override.find(qn("w:lvl")) is not None:
            return override.find(qn("w:lvl"))
        abstract_id = self._value(num, qn("w:abstractNumId"))
        abstract = self.root.find(f'w:abstractNum[@w:abstractNumId="{abstract_id}"]', self.root.nsmap)
        if abstract is None:
            return None
        return abstract.find(f'w:lvl[@w:ilvl="{level}"]', abstract.nsmap)

    @staticmethod
    def _format(value: int, fmt: str) -> str:
        if fmt in ("lowerLetter", "upperLetter") and value > 0:
            letters = ""
            while value:
                value, rem = divmod(value - 1, 26)
                letters = chr(65 + rem) + letters
            return letters.lower() if fmt == "lowerLetter" else letters
        if fmt in ("lowerRoman", "upperRoman") and 0 < value < 4000:
            roman = ""
            for number, label in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
                count, value = divmod(value, number)
                roman += label * count
            return roman.lower() if fmt == "lowerRoman" else roman
        if fmt == "decimalZero":
            return f"{value:02d}"
        return str(value)

    def label(self, paragraph: Any) -> str:
        if self.root is None:
            return ""
        # Direct properties override style inheritance; numId=0 disables a list.
        properties = [paragraph._p.pPr]
        style = paragraph.style
        seen_styles = set()
        while style is not None and style.style_id not in seen_styles:
            seen_styles.add(style.style_id)
            properties.append(style.element.pPr)
            style = style.base_style
        num_id = None
        level = None
        for props in properties:
            num_props = props.find(qn("w:numPr")) if props is not None else None
            if num_props is not None:
                if num_id is None:
                    num_id = self._value(num_props, qn("w:numId")) or None
                if level is None:
                    level = self._value(num_props, qn("w:ilvl")) or None
        if num_id in (None, "0"):
            return ""
        level = int(level or "0")
        num = self.root.find(f'w:num[@w:numId="{num_id}"]', self.root.nsmap)
        if num is None:
            return ""
        definition = self._level(num, level)
        if definition is None:
            return ""
        values = self.counters.setdefault(num_id, {})
        start = int(self._value(definition, qn("w:start"), "1"))
        override = num.find(f'w:lvlOverride[@w:ilvl="{level}"]', num.nsmap)
        start = int(self._value(override, qn("w:startOverride"), str(start)))
        values[level] = values[level] + 1 if level in values else start
        # By default a deeper level restarts after its parent. lvlRestart=0
        # explicitly disables that behavior; other values are one-based levels.
        for deeper in list(values):
            if deeper > level:
                restart = int(self._value(self._level(num, deeper), qn("w:lvlRestart"), str(deeper)))
                if restart and level < restart:
                    del values[deeper]
        fmt = self._value(definition, qn("w:numFmt"), "decimal")
        if fmt == "none":
            return ""
        if fmt == "bullet":
            return "- "
        template = self._value(definition, qn("w:lvlText"))
        if not template:
            return ""

        def substitute(match: Any) -> str:
            target = int(match.group(1)) - 1
            target_def = self._level(num, target)
            target_value = values.get(target, int(self._value(target_def, qn("w:start"), "1")))
            target_fmt = self._value(target_def, qn("w:numFmt"), "decimal")
            if definition.find(qn("w:isLgl")) is not None:
                target_fmt = "decimal"
            return self._format(target_value, target_fmt)

        return re.sub(r"%([1-9])", substitute, template) + " "


def _cell_content(cell: Any, numbering: Optional[_Numbering]) -> str:
    """Flatten cell blocks in their source order, retaining nested tables/images."""
    parts = []
    for child in cell._tc:
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, cell)
            text = paragraph.text.strip()
            # Even empty list paragraphs consume a source number in Word.
            label = numbering.label(paragraph) if numbering else ""
            if text:
                text = label + text
                parts.append(text)
            pictures = _image_placeholders(child)
            if pictures:
                parts.append((label if not text else "") + pictures)
        elif isinstance(child, CT_Tbl):
            nested = _table_to_markdown(Table(child, cell), numbering)
            if nested:
                parts.append(nested)
    return "\n".join(parts)


def _table_to_markdown(table: Any, numbering: Optional[_Numbering] = None) -> str:
    """Keep the table grid and emit each merged XML cell exactly once."""
    rows_data: List[List[str]] = []
    seen_cells = set()
    for row in table.rows:
        # Omitted cells still occupy grid positions in a Word table.
        cells = [""] * getattr(row, "grid_cols_before", 0)
        for cell in row.cells:
            if cell._tc in seen_cells:
                cells.append("")
                continue
            seen_cells.add(cell._tc)
            text = _cell_content(cell, numbering)
            cells.append(text.replace("\r\n", "\n").replace("\n", "<br>").replace("|", "\\|").strip())
        cells.extend([""] * getattr(row, "grid_cols_after", 0))
        rows_data.append(cells)
    max_cols = max((len(row) for row in rows_data), default=0)
    if not max_cols:
        return ""
    lines = []
    for index, row in enumerate(rows_data):
        row.extend([""] * (max_cols - len(row)))
        lines.append("| " + " | ".join(row) + " |")
        if index == 0:
            lines.append("| " + " | ".join(["---"] * max_cols) + " |")
    return "\n".join(lines)


def _image_placeholders(element: Any) -> str:
    """`<!--img:rIdN-->` for each picture in the element, so the ingest step can
    put the extracted image back at its original position."""
    try:
        rids = element.xpath(".//a:blip/@r:embed")
    except Exception:
        return ""
    return " ".join(image_placeholder(rid) for rid in rids)


def _has_page_break(paragraph_element: Any) -> bool:
    """Check if a paragraph element contains a page break."""
    xml_str = paragraph_element.xml
    return 'w:type="page"' in xml_str or "w:lastRenderedPageBreak" in xml_str


class DocxParser(BaseParser):
    """Parser for DOCX files preserving paragraph hierarchy, tables, and logical pages."""

    async def parse(
        self,
        file_data: bytes,
        file_name: str,
        *,
        vision_provider: Optional[Any] = None,
        tracker: Optional[Any] = None,
        **kwargs: Any,
    ) -> List[Dict[str, Any]]:
        """Parse DOCX binary content into Markdown pages."""
        if docx is None:
            raise RuntimeError("python-docx is not installed in the environment")

        try:
            doc = docx.Document(io.BytesIO(file_data))
        except Exception as e:
            logger.warning(f"DocxParser: python-docx failed to open '{file_name}': {e}. Falling back to mammoth.")
            return await self._fallback_mammoth(file_data, file_name)

        numbering = _Numbering(doc)
        pages: List[str] = []
        current_page_lines: List[str] = []

        # Iterate through document body in exact order
        for child in doc.element.body:
            if isinstance(child, CT_P):
                p = Paragraph(child, doc)
                if p.paragraph_format.page_break_before and current_page_lines:
                    pages.append("\n\n".join(current_page_lines))
                    current_page_lines = []
                text = p.text.strip()
                # Consume numbering before skipping empty paragraphs, so later
                # source labels do not shift after an empty list item.
                label = numbering.label(p)
                pics = _image_placeholders(child)
                if not text and pics:
                    current_page_lines.append(label + pics)
                    pics = ""
                if not text:
                    if _has_page_break(child):
                        if current_page_lines:
                            pages.append("\n\n".join(current_page_lines))
                            current_page_lines = []
                    continue

                text = label + text

                # Match actual heading styles, rather than substring matches
                # such as Subtitle. Markdown supports six heading levels.
                style_name = p.style.name.lower() if p.style and p.style.name else ""
                heading = re.fullmatch(r"heading\s+([1-9])", style_name)
                if heading:
                    text = f"{'#' * min(int(heading.group(1)), 6)} {text}"
                elif style_name == "title":
                    text = f"# {text}"

                current_page_lines.append(text)
                if pics:
                    current_page_lines.append(pics)

                if _has_page_break(child):
                    pages.append("\n\n".join(current_page_lines))
                    current_page_lines = []

            elif isinstance(child, CT_Tbl):
                tbl = Table(child, doc)
                tbl_md = _table_to_markdown(tbl, numbering)
                if tbl_md:
                    current_page_lines.append(tbl_md)

        if current_page_lines:
            pages.append("\n\n".join(current_page_lines))

        if not pages:
            pages = [""]

        # Post-process each page with Vietnamese legal headings and NFC normalizer
        final_pages: List[Dict[str, Any]] = []
        for idx, page_content in enumerate(pages):
            enhanced = enhance_vietnamese_headings(page_content)
            normalized = normalize_text(enhanced)
            final_pages.append({
                "content": normalized,
                "page_number": idx + 1,
            })

        return final_pages

    async def _fallback_mammoth(self, file_data: bytes, file_name: str) -> List[Dict[str, Any]]:
        """Fallback to mammoth markdown conversion if python-docx fails."""
        try:
            import mammoth
            result = mammoth.convert_to_markdown(io.BytesIO(file_data))
            text = normalize_text(enhance_vietnamese_headings(result.value or ""))
            return [{"content": text, "page_number": 1}]
        except Exception as err:
            logger.error(f"DocxParser: mammoth fallback also failed for '{file_name}': {err}")
            return [{"content": "", "page_number": 1}]

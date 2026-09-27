"""PPTX presentation parser for structured markdown extraction."""

import io
from typing import Any, Dict, List, Optional
from loguru import logger
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.parsers.base import BaseParser


class PPTXParser(BaseParser):
    """Parser for PowerPoint presentations (.pptx).
    
    Converts presentations into page-per-slide structured markdown,
    extracting titles, bullet points, tables, and speaker notes.
    """

    def _extract_table_markdown(self, table) -> str:
        """Convert a pptx table shape into markdown table syntax."""
        rows = list(table.rows)
        if not rows:
            return ""

        grid = []
        for row in rows:
            row_cells = [cell.text.replace("\n", " ").replace("\r", " ").strip() for cell in row.cells]
            grid.append(row_cells)

        if not grid:
            return ""

        col_count = len(grid[0])
        # Ensure all rows have the same number of columns
        for row in grid:
            while len(row) < col_count:
                row.append("")

        header = "| " + " | ".join(grid[0]) + " |"
        sep = "| " + " | ".join(["---"] * col_count) + " |"
        body_lines = ["| " + " | ".join(row) + " |" for row in grid[1:]]

        lines = [header, sep] + body_lines
        return "\n".join(lines)

    def _extract_shape_text(self, shape, title_id: Optional[int]) -> List[str]:
        """Extract text or markdown from a shape, including nested group shapes."""
        chunks: List[str] = []

        # Recurse into group shapes
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP or hasattr(shape, "shapes"):
            for child in shape.shapes:
                chunks.extend(self._extract_shape_text(child, title_id))
            return chunks

        # Skip title shape if already extracted
        if title_id is not None and shape.shape_id == title_id:
            return chunks

        # Extract table
        if shape.has_table:
            tbl_md = self._extract_table_markdown(shape.table)
            if tbl_md:
                chunks.append(tbl_md)
            return chunks

        # Extract text frame
        if shape.has_text_frame:
            tf = shape.text_frame
            para_texts = []
            for p in tf.paragraphs:
                txt = p.text.strip()
                if not txt:
                    continue
                indent = "  " * getattr(p, "level", 0)
                if getattr(p, "level", 0) > 0:
                    para_texts.append(f"{indent}- {txt}")
                else:
                    para_texts.append(txt)
            if para_texts:
                chunks.append("\n".join(para_texts))

        return chunks

    async def parse(
        self,
        file_data: bytes,
        file_name: str,
        *,
        vision_provider: Optional[Any] = None,
        tracker: Optional[Any] = None,
        **kwargs: Any,
    ) -> List[Dict[str, Any]]:
        """Extract structured text from binary PPTX file data.

        Returns:
            List of page records: [{"content": str, "page_number": int}]
        """
        try:
            prs = Presentation(io.BytesIO(file_data))
        except Exception as e:
            logger.warning(f"PPTXParser: Failed to open PPTX presentation '{file_name}': {e}")
            raise ValueError(f"Không thể giải mã file PPTX '{file_name}': {e}") from e

        slides = list(prs.slides)
        if not slides:
            logger.info(f"PPTXParser: Presentation '{file_name}' contains no slides.")
            return []

        pages: List[Dict[str, Any]] = []

        for slide_idx, slide in enumerate(slides, 1):
            slide_parts: List[str] = []

            # 1. Slide Title
            title = ""
            title_id = None
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text.strip()
                title_id = slide.shapes.title.shape_id

            if title:
                slide_parts.append(f"# Slide {slide_idx}: {title}")
            else:
                slide_parts.append(f"# Slide {slide_idx}")

            # 2. Slide Shapes & Content
            body_parts: List[str] = []
            for shape in slide.shapes:
                extracted = self._extract_shape_text(shape, title_id)
                body_parts.extend(extracted)

            if body_parts:
                slide_parts.append("\n\n".join(body_parts))

            # 3. Speaker Notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_txt = slide.notes_slide.notes_text_frame.text.strip()
                if notes_txt:
                    slide_parts.append(f"> **Ghi chú diễn giả (Speaker Notes):**\n> {notes_txt}")

            content_md = "\n\n".join(slide_parts).strip()
            pages.append({
                "page_number": slide_idx,
                "content": content_md,
            })

        logger.info(f"PPTXParser: Extracted {len(pages)} slides from '{file_name}'.")
        return pages

"""Image extraction from PDF/DOCX/PPTX and image-derived retrieval chunks."""

from __future__ import annotations

import io
from typing import Any, Optional

from loguru import logger

from app.ai.embedding_catalog import get_spec
from app.ai.registry import ProviderRegistry
from app.services.storage_service import storage_service

try:  # provided by the retrieval domain
    from app.services.embedding_storage import chunk_content_hash, upsert_chunk_embedding
except ImportError:  # pragma: no cover - resolved lazily until the module lands

    async def upsert_chunk_embedding(*args: Any, **kwargs: Any) -> Any:
        from app.services import embedding_storage

        return await embedding_storage.upsert_chunk_embedding(*args, **kwargs)

    def chunk_content_hash(text: str) -> str:
        from app.services import embedding_storage

        return embedding_storage.chunk_content_hash(text)


MIN_IMAGE_BYTES = 2048
IMAGE_CHUNK_INDEX_BASE = 80000

MIME_BY_EXT = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "gif": "image/gif",
    "bmp": "image/bmp",
    "tiff": "image/tiff",
    "webp": "image/webp",
    "svg": "image/svg+xml",
}


def _mime_for(ext: str) -> str:
    return MIME_BY_EXT.get((ext or "").lower(), "image/png")


class ImageInfo:
    def __init__(
        self,
        minio_key: str,
        page_number: Optional[int],
        image_index: int,
        content_type: str,
        size_bytes: int,
        caption: Optional[str] = None,
        image_id: Optional[str] = None,
        ref: Optional[str] = None,
    ):
        self.minio_key = minio_key
        self.page_number = page_number
        self.image_index = image_index
        self.content_type = content_type
        self.size_bytes = size_bytes
        self.caption = caption
        self.image_id = image_id
        self.ref = ref

    def __repr__(self) -> str:
        return f"ImageInfo(key={self.minio_key!r}, page={self.page_number}, index={self.image_index})"


def _store(key: str, data: bytes, content_type: str) -> None:
    storage_service.upload_file(object_name=key, data=data, content_type=content_type)


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------


def extract_images_from_pdf(file_data: bytes, source_id: Any) -> list[ImageInfo]:
    try:
        import fitz

        doc = fitz.open(stream=file_data, filetype="pdf")
    except Exception as e:  # noqa: BLE001
        logger.warning(f"extract_images_from_pdf: cannot open PDF for {source_id}: {e}")
        return []
    images: list[ImageInfo] = []
    index = 0
    try:
        for page_idx in range(len(doc)):
            page_no = page_idx + 1
            try:
                refs = doc[page_idx].get_images(full=True)
            except Exception as e:  # noqa: BLE001
                logger.warning(f"extract_images_from_pdf: page {page_no} listing failed: {e}")
                continue
            for ref in refs:
                try:
                    info = doc.extract_image(ref[0])
                    data = info.get("image") or b""
                    if len(data) < MIN_IMAGE_BYTES:
                        continue
                    ext = (info.get("ext") or "png").lower()
                    content_type = _mime_for(ext)
                    key = f"sources/{source_id}/images/page{page_no}_{index}.{ext}"
                    _store(key, data, content_type)
                    images.append(ImageInfo(key, page_no, index, content_type, len(data)))
                    index += 1
                except Exception as e:  # noqa: BLE001
                    logger.warning(f"extract_images_from_pdf: skipping image on page {page_no}: {e}")
    finally:
        try:
            doc.close()
        except Exception:  # noqa: BLE001
            pass
    return images


def _docx_image_parts(file_data: bytes):
    import docx

    document = docx.Document(io.BytesIO(file_data))
    for rel_id, rel in document.part.rels.items():
        if "image" not in (rel.reltype or ""):
            continue
        if getattr(rel, "is_external", False):
            continue
        yield rel_id, rel.target_part


def extract_images_from_docx(file_data: bytes, source_id: Any) -> list[ImageInfo]:
    images: list[ImageInfo] = []
    index = 0
    try:
        parts = list(_docx_image_parts(file_data))
    except Exception as e:  # noqa: BLE001
        logger.warning(f"extract_images_from_docx: cannot open DOCX for {source_id}: {e}")
        return []
    for rel_id, part in parts:
        try:
            data = part.blob
            if len(data) < MIN_IMAGE_BYTES:
                continue
            content_type = getattr(part, "content_type", None) or "image/png"
            ext = content_type.split("/")[-1].lower()
            if ext == "svg+xml":
                ext = "svg"
            key = f"sources/{source_id}/images/docx_{index}.{ext}"
            _store(key, data, content_type)
            images.append(ImageInfo(key, None, index, content_type, len(data), ref=rel_id))
            index += 1
        except Exception as e:  # noqa: BLE001
            logger.warning(f"extract_images_from_docx: skipping {rel_id}: {e}")
    return images


def docx_image_blobs(file_data: bytes) -> dict[str, tuple[str, bytes]]:
    try:
        return {
            rel_id: (getattr(part, "content_type", None) or "image/png", part.blob)
            for rel_id, part in _docx_image_parts(file_data)
        }
    except Exception as e:  # noqa: BLE001
        logger.warning(f"docx_image_blobs: cannot open DOCX: {e}")
        return {}


def _picture_shapes(shapes) -> list:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    found = []
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                found.extend(_picture_shapes(shape.shapes))
                continue
        except Exception:  # noqa: BLE001
            pass
        if hasattr(shape, "image"):
            try:
                shape.image  # noqa: B018 - placeholders without an image raise here
            except Exception:  # noqa: BLE001
                continue
            found.append(shape)
    return found


def extract_images_from_pptx(file_data: bytes, source_id: Any) -> list[ImageInfo]:
    try:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(file_data))
    except Exception as e:  # noqa: BLE001
        logger.warning(f"extract_images_from_pptx: cannot open PPTX for {source_id}: {e}")
        return []
    images: list[ImageInfo] = []
    index = 0
    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape in _picture_shapes(slide.shapes):
            try:
                image = shape.image
                data = image.blob
                ext = (image.ext or "png").lower()
                content_type = image.content_type or _mime_for(ext)
                key = f"sources/{source_id}/images/slide{slide_no}_{index}.{ext}"
                _store(key, data, content_type)
                images.append(ImageInfo(key, slide_no, index, content_type, len(data)))
                index += 1
            except Exception as e:  # noqa: BLE001
                logger.warning(f"extract_images_from_pptx: skipping image on slide {slide_no}: {e}")
    return images


def extract_images(file_data: bytes, file_name: str, source_id: Any) -> list[ImageInfo]:
    ext = (file_name or "").rsplit(".", 1)[-1].lower() if "." in (file_name or "") else ""
    if ext == "pdf":
        return extract_images_from_pdf(file_data, source_id)
    if ext == "docx":
        return extract_images_from_docx(file_data, source_id)
    if ext == "pptx":
        return extract_images_from_pptx(file_data, source_id)
    return []


# ---------------------------------------------------------------------------
# Image chunks
# ---------------------------------------------------------------------------


def create_image_chunk(image_info: ImageInfo, chunk_type: str, text: str) -> dict:
    page = image_info.page_number or 1
    body = text or ""
    if not body.strip().startswith(f"[{chunk_type}"):
        body = f"[{chunk_type} | Trang {page}]: {body}"
    return {
        "chunk_type": chunk_type,
        "text": body,
        "page_number": page,
        "image_id": image_info.image_id,
        "minio_key": image_info.minio_key,
    }


async def index_image_chunks(session: Any, source_id: Any, chunks: list[dict], spec_id: Optional[str] = None) -> int:
    if not chunks:
        return 0
    registry = ProviderRegistry(session)
    if spec_id is None:
        spec_id = await registry.get_active_embedding_spec_id()
    if not spec_id:
        return 0
    spec = get_spec(spec_id)
    provider = await registry.get_embedding(task="document", spec_id=spec_id)
    texts = [c.get("text") or "" for c in chunks]
    vectors = await provider.embed_batch(texts)
    for i, (chunk, vector) in enumerate(zip(chunks, vectors)):
        text = texts[i]
        await upsert_chunk_embedding(
            session=session,
            source_id=source_id,
            chunk_index=IMAGE_CHUNK_INDEX_BASE + i,
            spec=spec,
            vector=vector,
            text=text,
            start_char=0,
            end_char=len(text),
            page_number=chunk.get("page_number") or 1,
            content_hash=chunk_content_hash(text),
        )
    return len(chunks)

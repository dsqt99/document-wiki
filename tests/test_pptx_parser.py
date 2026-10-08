"""Tests for PPTX presentation parsing, image extraction, and visual chunking."""

import io
import os
import uuid
import pytest
from unittest.mock import AsyncMock, MagicMock, patch
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.services.parsers.pptx_parser import PPTXParser
from app.services.image_service import (
    extract_images,
    extract_images_from_pptx,
    create_image_chunk,
    index_image_chunks,
    ImageInfo,
)


def _create_sample_image_bytes(size: int = 100) -> bytes:
    """Create a sample PNG image large enough to pass MIN_IMAGE_BYTES."""
    img = Image.frombytes("RGB", (size, size), os.urandom(size * size * 3))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _build_sample_pptx(include_image: bool = False) -> bytes:
    """Build an in-memory PPTX presentation with titles, text, tables, and notes."""
    prs = Presentation()

    # Slide 1: Title + Body
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Kế hoạch Tuần tra Kiểm soát"
    body1 = slide1.placeholders[1]
    body1.text = "Mục tiêu đảm bảo ANTT năm 2026\nĐội Cảnh sát Giao thông"

    # Slide 2: Table + Speaker Notes
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    # Add table
    table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Tuyến đường"
    table.cell(0, 1).text = "Quân số"
    table.cell(1, 0).text = "Quốc lộ 5"
    table.cell(1, 1).text = "12 đồng chí"

    # Add speaker notes
    notes_slide = slide2.notes_slide
    notes_slide.notes_text_frame.text = "Lưu ý kiểm tra nồng độ cồn vào ban đêm."

    # Slide 3: Image (optional)
    if include_image:
        slide3 = prs.slides.add_slide(prs.slide_layouts[6])
        img_bytes = _create_sample_image_bytes()
        slide3.shapes.add_picture(io.BytesIO(img_bytes), Inches(1), Inches(1), Inches(3), Inches(3))

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_parse_pptx_basic():
    """Test extracting slides, markdown tables, and speaker notes from PPTX."""
    pptx_bytes = _build_sample_pptx(include_image=False)
    parser = PPTXParser()

    pages = await parser.parse(pptx_bytes, "ke_hoach.pptx")

    assert len(pages) == 2
    # Slide 1 check
    page1 = pages[0]
    assert page1["page_number"] == 1
    assert "Kế hoạch Tuần tra Kiểm soát" in page1["content"]
    assert "Mục tiêu đảm bảo ANTT" in page1["content"]

    # Slide 2 check
    page2 = pages[1]
    assert page2["page_number"] == 2
    assert "Tuyến đường" in page2["content"]
    assert "Quốc lộ 5" in page2["content"]
    assert "12 đồng chí" in page2["content"]
    assert "Ghi chú diễn giả" in page2["content"] or "Speaker Notes" in page2["content"]
    assert "Lưu ý kiểm tra nồng độ cồn" in page2["content"]


@pytest.mark.asyncio
async def test_parse_pptx_empty_and_corrupted():
    """Test handling empty presentation and corrupted bytes."""
    parser = PPTXParser()

    # Empty presentation (0 slides)
    empty_prs = Presentation()
    empty_buf = io.BytesIO()
    empty_prs.save(empty_buf)
    pages = await parser.parse(empty_buf.getvalue(), "empty.pptx")
    assert pages == []

    # Corrupted bytes
    with pytest.raises(ValueError, match="Không thể giải mã file PPTX"):
        await parser.parse(b"NOT_A_VALID_PPTX_FILE", "corrupted.pptx")


@pytest.mark.asyncio
async def test_extract_images_from_pptx():
    """Test extracting pictures from PPTX presentation with page numbers."""
    pptx_bytes = _build_sample_pptx(include_image=True)
    source_id = str(uuid.uuid4())

    with patch("app.services.image_service.storage_service") as mock_storage:
        mock_storage.upload_file.return_value = True

        images = extract_images_from_pptx(pptx_bytes, source_id)
        assert len(images) == 1
        img = images[0]
        assert img.page_number == 3
        assert img.content_type == "image/png"
        assert img.size_bytes > 0
        assert f"sources/{source_id}/images/slide3_0.png" in img.minio_key
        mock_storage.upload_file.assert_called_once()


@pytest.mark.asyncio
async def test_extract_images_dispatcher_with_pptx():
    """Test extract_images auto-detecting pptx extension."""
    pptx_bytes = _build_sample_pptx(include_image=True)
    source_id = str(uuid.uuid4())

    with patch("app.services.image_service.storage_service") as mock_storage:
        mock_storage.upload_file.return_value = True
        images = extract_images(pptx_bytes, "presentation.pptx", source_id)
        assert len(images) == 1
        assert images[0].page_number == 3


def test_create_image_chunk():
    """Test creating structured image_caption and image_ocr chunks."""
    img_info = ImageInfo(
        minio_key="sources/123/images/slide1_0.png",
        page_number=2,
        image_index=0,
        content_type="image/png",
        size_bytes=4096,
        caption="Biểu đồ thống kê tội phạm năm 2025",
        image_id=str(uuid.uuid4()),
    )

    caption_chunk = create_image_chunk(
        img_info,
        chunk_type="image_caption",
        text=img_info.caption,
    )
    assert caption_chunk["chunk_type"] == "image_caption"
    assert caption_chunk["page_number"] == 2
    assert "Biểu đồ thống kê tội phạm" in caption_chunk["text"]
    assert caption_chunk["image_id"] == img_info.image_id

    ocr_chunk = create_image_chunk(
        img_info,
        chunk_type="image_ocr",
        text="Công an tỉnh Hưng Yên - Phòng Cảnh sát ĐTTP",
    )
    assert ocr_chunk["chunk_type"] == "image_ocr"
    assert ocr_chunk["page_number"] == 2
    assert "Phòng Cảnh sát ĐTTP" in ocr_chunk["text"]


@pytest.mark.asyncio
async def test_index_image_chunks_with_mock_db():
    """Test indexing image chunks into SourceChunkEmbedding."""
    session = AsyncMock()
    source_id = uuid.uuid4()
    chunks = [
        {
            "chunk_type": "image_caption",
            "text": "[image_caption | Trang 1]: Sơ đồ tổ chức cơ quan",
            "page_number": 1,
            "image_id": str(uuid.uuid4()),
            "minio_key": "sources/test/img1.png",
        }
    ]

    mock_spec = MagicMock()
    mock_spec.id = "text-embedding-3-small"
    mock_spec.dimension = 768

    mock_provider = AsyncMock()
    mock_provider.embed_batch.return_value = [[0.1] * 768]

    with patch("app.services.image_service.ProviderRegistry") as mock_reg_cls, \
         patch("app.services.image_service.get_spec", return_value=mock_spec), \
         patch("app.services.image_service.upsert_chunk_embedding", new_callable=AsyncMock) as mock_upsert:

        mock_registry = MagicMock()
        mock_registry.get_active_embedding_spec_id = AsyncMock(return_value="text-embedding-3-small")
        mock_registry.get_embedding = AsyncMock(return_value=mock_provider)
        mock_reg_cls.return_value = mock_registry

        indexed_count = await index_image_chunks(session, source_id, chunks)

        assert indexed_count == 1
        mock_provider.embed_batch.assert_called_once()
        mock_upsert.assert_called_once()
        call_kwargs = mock_upsert.call_args[1]
        assert call_kwargs["source_id"] == source_id
        assert call_kwargs["page_number"] == 1
        assert "Sơ đồ tổ chức cơ quan" in call_kwargs["text"]

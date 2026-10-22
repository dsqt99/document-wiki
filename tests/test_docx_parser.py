import io
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.services.parsers.docx_parser import DocxParser


def create_sample_docx_with_table() -> bytes:
    """Create an in-memory docx with headings, paragraphs, and a markdown table."""
    doc = Document()
    doc.add_heading("QUY CHẾ HOẠT ĐỘNG", level=1)
    doc.add_paragraph("Điều 1. Phạm vi điều chỉnh")
    doc.add_paragraph("Nội dung quy định chi tiết.")

    # Add a table 2x2
    table = doc.add_table(rows=2, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Họ và tên"
    hdr_cells[1].text = "Chức vụ"

    row_cells = table.rows[1].cells
    row_cells[0].text = "Nguyễn Văn A"
    row_cells[1].text = "Cán bộ"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_docx_with_merged_cells() -> bytes:
    """Create an in-memory docx with horizontally merged cells."""
    doc = Document()
    doc.add_paragraph("Bảng thống kê")

    table = doc.add_table(rows=3, cols=3)
    # Merge row 0 cols 1 and 2
    cell_a = table.cell(0, 1)
    cell_b = table.cell(0, 2)
    cell_a.merge(cell_b)

    table.cell(0, 0).text = "STT"
    table.cell(0, 1).text = "Thông tin chung"
    
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "Tên: Nguyễn Văn B"
    table.cell(1, 2).text = "Tuổi: 30"

    table.cell(2, 0).text = "2"
    table.cell(2, 1).text = "Tên: Trần Thị C"
    table.cell(2, 2).text = "Tuổi: 25"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_docx_parser_extracts_headings_and_tables():
    """Verify DocxParser extracts headings and formats tables into clean Markdown."""
    docx_bytes = create_sample_docx_with_table()
    parser = DocxParser()
    pages = await parser.parse(docx_bytes, "sample.docx")

    assert len(pages) >= 1
    content = pages[0]["content"]
    assert "QUY CHẾ HOẠT ĐỘNG" in content
    assert "### Điều 1. Phạm vi điều chỉnh" in content
    # Check table format
    assert "| Họ và tên | Chức vụ |" in content
    assert "| Nguyễn Văn A | Cán bộ |" in content


@pytest.mark.asyncio
async def test_docx_parser_handles_merged_cells():
    """Verify DocxParser preserves table structure without dropping columns when cells are merged."""
    docx_bytes = create_sample_docx_with_merged_cells()
    parser = DocxParser()
    pages = await parser.parse(docx_bytes, "merged.docx")

    assert len(pages) >= 1
    content = pages[0]["content"]
    assert "| STT |" in content
    assert "Nguyễn Văn B" in content
    assert "Trần Thị C" in content


def _docx_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_docx_preserves_paragraph_table_order_and_nested_cell_content():
    doc = Document()
    doc.add_paragraph("Before table")
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "Before nested"
    nested = cell.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Nested first"
    nested.cell(0, 1).text = "Nested second"
    cell.add_paragraph("After nested")
    doc.add_paragraph("After table")
    content = (await DocxParser().parse(_docx_bytes(doc), "nested.docx"))[0]["content"]
    labels = ["Before table", "Before nested", "Nested first", "Nested second", "After nested", "After table"]
    assert all(label in content for label in labels)
    assert [content.index(label) for label in labels] == sorted(content.index(label) for label in labels)


@pytest.mark.asyncio
async def test_docx_merged_cells_emit_content_once_without_removing_equal_independent_cells():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Merged vertical"
    table.cell(1, 1).text = table.cell(1, 2).text = "Equal independent"
    content = (await DocxParser().parse(_docx_bytes(doc), "merged.docx"))[0]["content"]
    assert content.count("Merged heading") == 1
    assert content.count("Merged vertical") == 1
    assert content.count("Equal independent") == 2
    assert "| Merged heading | | |" in content


@pytest.mark.asyncio
async def test_docx_restores_style_numbering_across_tables():
    doc = Document()
    doc.add_paragraph("First item", style="List Number")
    doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].text = "Between items"
    doc.add_paragraph("Second item", style="List Number")
    content = (await DocxParser().parse(_docx_bytes(doc), "numbering.docx"))[0]["content"]
    assert "1. First item" in content
    assert "2. Second item" in content


def _set_numbering(doc, paragraph, *, num_id, abstract_id, levels, overrides=""):
    from docx.oxml.ns import qn
    numbering = doc.part.numbering_part.element
    if not numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]'):
        numbering.append(parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">{levels}</w:abstractNum>'))
    if not numbering.xpath(f'./w:num[@w:numId="{num_id}"]'):
        numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{overrides}</w:num>'))
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_numId().val = num_id
    return num_pr


@pytest.mark.asyncio
async def test_docx_restores_legal_numbering_level_formats_and_restart_overrides():
    doc = Document()
    levels = '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="Điều %1."/></w:lvl>'
    levels += '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/></w:lvl>'
    for text, level in [("Scope", 0), ("First point", 1), ("Second point", 1), ("Other scope", 0), ("Restart point", 1)]:
        p = doc.add_paragraph(text)
        num_pr = _set_numbering(doc, p, num_id=20, abstract_id=20, levels=levels, overrides='<w:lvlOverride w:ilvl="0"><w:startOverride w:val="4"/></w:lvlOverride>')
        num_pr.get_or_add_ilvl().val = level
    content = (await DocxParser().parse(_docx_bytes(doc), "legal-numbering.docx"))[0]["content"]
    for text in ["Điều 4. Scope", "a) First point", "b) Second point", "Điều 5. Other scope", "a) Restart point"]:
        assert text in content


@pytest.mark.asyncio
async def test_docx_explicit_zero_numid_disables_inherited_numbering():
    doc = Document()
    p = doc.add_paragraph("Not numbered", style="List Number")
    p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
    content = (await DocxParser().parse(_docx_bytes(doc), "disabled-list.docx"))[0]["content"]
    assert content == "Not numbered"


@pytest.mark.asyncio
async def test_docx_preserves_heading_levels_and_does_not_promote_subtitle():
    doc = Document()
    doc.add_heading("Fourth heading", level=4)
    doc.add_paragraph("Document subtitle", style="Subtitle")
    content = (await DocxParser().parse(_docx_bytes(doc), "styles.docx"))[0]["content"]
    assert "#### Fourth heading" in content
    assert "# Document subtitle" not in content


@pytest.mark.asyncio
async def test_docx_numbering_preserves_roman_multilevel_and_never_restart():
    doc = Document()
    levels = '<w:lvl w:ilvl="0"><w:start w:val="3"/><w:numFmt w:val="upperRoman"/><w:lvlText w:val="%1."/></w:lvl>'
    levels += '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlRestart w:val="0"/><w:lvlText w:val="%1.%2"/></w:lvl>'
    for text, level in [("Parent", 0), ("Child", 1), ("Next parent", 0), ("Next child", 1)]:
        p = doc.add_paragraph(text)
        _set_numbering(doc, p, num_id=25, abstract_id=25, levels=levels).get_or_add_ilvl().val = level
    content = (await DocxParser().parse(_docx_bytes(doc), "levels.docx"))[0]["content"]
    for text in ["III. Parent", "III.1 Child", "IV. Next parent", "IV.2 Next child"]:
        assert text in content


@pytest.mark.asyncio
async def test_docx_paragraph_page_break_before_starts_next_page():
    doc = Document()
    doc.add_paragraph("First page")
    p = doc.add_paragraph("Second page")
    p.paragraph_format.page_break_before = True
    pages = await DocxParser().parse(_docx_bytes(doc), "page-before.docx")
    assert [page["content"] for page in pages] == ["First page", "Second page"]


@pytest.mark.asyncio
async def test_docx_without_optional_numbering_part_still_extracts_text():
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc = Document()
    doc.add_paragraph("Plain document without lists")
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.NUMBERING:
            doc.part.drop_rel(rel_id)
    content = (await DocxParser().parse(_docx_bytes(doc), "no-numbering.docx"))[0]["content"]
    assert content == "Plain document without lists"


@pytest.mark.asyncio
async def test_docx_empty_numbered_paragraphs_advance_source_labels_without_emitting_text():
    doc = Document()
    doc.add_paragraph("First item", style="List Number")
    doc.add_paragraph("", style="List Number")
    doc.add_paragraph("")
    doc.add_paragraph("Third item", style="List Number")
    content = (await DocxParser().parse(_docx_bytes(doc), "empty-list.docx"))[0]["content"]
    assert "1. First item" in content
    assert "3. Third item" in content
    assert "2." not in content


@pytest.mark.asyncio
async def test_docx_empty_numbered_table_cell_paragraphs_advance_source_labels():
    doc = Document()
    doc.add_paragraph("First item", style="List Number")
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].style = doc.styles["List Number"]
    cell.add_paragraph("")
    doc.add_paragraph("Third item", style="List Number")
    content = (await DocxParser().parse(_docx_bytes(doc), "empty-cell-list.docx"))[0]["content"]
    assert "1. First item" in content
    assert "3. Third item" in content
    assert "2." not in content
